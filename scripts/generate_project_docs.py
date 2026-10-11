from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
    KeepTogether, HRFlowable,
)
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PDF_PATH = ROOT / 'FoodRescue_Project_Overview.pdf'
PPT_PATH = ROOT / 'FoodRescue_Project_Overview.pptx'

NAVY = colors.HexColor('#10243A')
GREEN = colors.HexColor('#2B8A67')
GOLD = colors.HexColor('#F3B61F')
PALE = colors.HexColor('#F3F7F5')
INK = colors.HexColor('#243447')
MUTED = colors.HexColor('#5C6B78')
WHITE = colors.white


def build_pdf():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CoverTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=31, leading=37, textColor=NAVY, alignment=TA_LEFT, spaceAfter=8))
    styles.add(ParagraphStyle(name='CoverSub', parent=styles['Normal'], fontSize=14, leading=21, textColor=GREEN, spaceAfter=16))
    styles.add(ParagraphStyle(name='Section', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=NAVY, spaceBefore=8, spaceAfter=9, keepWithNext=True))
    styles.add(ParagraphStyle(name='Subsection', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=GREEN, spaceBefore=7, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name='BodyX', parent=styles['BodyText'], fontSize=9.2, leading=14, textColor=INK, spaceAfter=6))
    styles.add(ParagraphStyle(name='SmallX', parent=styles['BodyText'], fontSize=8, leading=11, textColor=MUTED, spaceAfter=3))
    styles.add(ParagraphStyle(name='BulletX', parent=styles['BodyText'], fontSize=9, leading=13, leftIndent=13, firstLineIndent=-8, bulletIndent=0, textColor=INK, spaceAfter=3))
    styles.add(ParagraphStyle(name='CellHead', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=8.3, leading=10, textColor=WHITE))
    styles.add(ParagraphStyle(name='CellX', parent=styles['BodyText'], fontSize=7.7, leading=10, textColor=INK))
    styles.add(ParagraphStyle(name='Callout', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=9, leading=13, textColor=NAVY, backColor=PALE, borderColor=GREEN, borderWidth=0.7, borderPadding=8, spaceBefore=4, spaceAfter=10))

    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=19*mm, bottomMargin=18*mm, title='FoodRescue Project Overview', author='FoodRescue Project')
    story = []
    P = lambda text, style='BodyX': Paragraph(text, styles[style])
    bullet = lambda text: Paragraph('• ' + text, styles['BulletX'])

    story += [Spacer(1, 30*mm), P('FoodRescue', 'CoverTitle'), P('Project Overview & Technical Guide', 'CoverSub'), HRFlowable(width='100%', thickness=2, color=GOLD, spaceAfter=15)]
    story += [P('<b>Purpose</b><br/>A food donation prototype connecting donors who have surplus food with NGOs that can respond to pickup requests.', 'BodyX'), Spacer(1, 8)]
    overview = [
        [P('<b>Frontend</b><br/>Next.js App Router · React · TypeScript', 'CellX'), P('<b>Backend</b><br/>Next.js route handlers · server-side validation', 'CellX')],
        [P('<b>Data</b><br/>Prisma ORM · SQLite local database', 'CellX'), P('<b>Assistant</b><br/>Google Gemini Developer API', 'CellX')],
    ]
    t = Table(overview, colWidths=[82*mm, 82*mm], rowHeights=[23*mm, 23*mm])
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),PALE),('BOX',(0,0),(-1,-1),0.6,GREEN),('INNERGRID',(0,0),(-1,-1),0.5,colors.HexColor('#D5E3DC')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10)]))
    story += [t, Spacer(1, 15), P('Prepared from the source code and configuration in this workspace. This document excludes secret values.', 'SmallX'), P('<b>Security:</b> Rotate any Gemini API key previously stored in the sample environment file. Keep replacement credentials only in a local ignored environment file.', 'Callout'), PageBreak()]

    story += [P('1. Product and User Flows', 'Section'), P('FoodRescue supports three roles: donors, NGO operators, and administrators. The product is a prototype; listings and actions are not proof of real-world verification or completed pickups.', 'BodyX'), P('Donor flow', 'Subsection')]
    for text in [
        'Enter pickup location, region, food type, positive quantity/unit, and food expiry.',
        'Optionally use browser location; the server calls OpenStreetMap Nominatim reverse geocoding to suggest a readable address and region.',
        'Review NGO choices at <b>/select-ngo</b>, then broadcast the donation. Broadcasts expire at the earlier of one hour or the food expiry (as described in the chatbot system prompt).',
        'Donation validation requires expiry at least ten minutes in the future; supported units are kg, servings, boxes, and litres.',
    ]: story.append(bullet(text))
    story += [P('NGO flow', 'Subsection')]
    for text in ['Register or sign in at <b>/login</b>. NGO accounts must be active and approved to authenticate.', 'View unexpired broadcasts for the assigned region; Unassigned NGOs can see requests across regions. Busy NGOs do not receive open requests.', 'Accept a request, decline it individually, or move an accepted pickup through <b>in_progress</b> and <b>completed</b>. The dashboard refreshes requests periodically.']: story.append(bullet(text))
    story += [P('Administrator flow', 'Subsection')]
    for text in ['Sign in at <b>/admin/login</b>. The admin portal provides regional summaries and NGO account management.', 'Review NGO accounts and manage region assignment/activation and approval state. Admin access is protected by a signed session cookie.']: story.append(bullet(text))

    story += [P('2. Frontend Architecture', 'Section')]
    front_rows = [
        ['Area', 'Routes / files', 'Responsibility'],
        ['/ landing', 'app/page.tsx; components/hero-section.tsx', 'Donor-facing entry point and impact content.'],
        ['NGO selection', 'app/select-ngo/page.tsx; components/ngo/*', 'Reads search parameters and builds NGO selection experience.'],
        ['NGO portal', 'app/login/page.tsx; app/ngo/dashboard/page.tsx; components/ngo-dashboard.tsx', 'Registration/login, request feed and pickup status actions.'],
        ['Admin portal', 'app/admin/login/page.tsx; app/admin/page.tsx; components/admin-dashboard.tsx', 'Admin authentication and operations dashboard.'],
        ['Global UI', 'app/layout.tsx; components/site-header.tsx; components/help-chatbot.tsx', 'Shared layout, navigation and global chat widget.'],
        ['Styling', 'app/globals.css; Tailwind CSS 4', 'Global theme and utility-based styling.'],
    ]
    story.append(make_pdf_table(front_rows, [27*mm, 70*mm, 75*mm], styles))
    story += [Spacer(1, 8), P('The app uses the Next.js App Router. Server components load protected pages/data; interactive widgets are client components. Alias <b>@/*</b> maps to the workspace root in tsconfig.json.', 'BodyX')]

    story += [P('3. Backend and API', 'Section'), P('API endpoints are Next.js route handlers under app/api. They parse and validate input, authenticate protected requests, update records through Prisma, and return JSON responses.', 'BodyX')]
    api_rows = [
        ['API area', 'Routes', 'Main behavior'],
        ['Donations', '/api/donations; /api/donations/[id]; /api/donations/[id]/status', 'Create donation broadcasts; retrieve/update donation data and status.'],
        ['NGO requests', '/api/ngo-requests; /[id]/accept; /[id]/decline; /[id]/status', 'List eligible requests, accept/decline, advance pickup status. Requires NGO session.'],
        ['NGO auth', '/api/ngo-auth/register, /login, /logout, /availability', 'Create account, authenticate, end session, update availability.'],
        ['Admin', '/api/admin/dashboard; /api/admin/ngos; /api/admin/ngos/[id]', 'Dashboard summaries and NGO account management. Admin protected.'],
        ['Admin auth', '/api/admin-auth/login; /api/admin-auth/logout', 'Admin sign-in and sign-out.'],
        ['Assistant', 'POST /api/help', 'Validates conversation, enforces in-memory rate limit, calls Gemini.'],
        ['Geocoding', 'GET /api/location/reverse?lat=…&lng=…', 'Reverse lookup using Nominatim; returns address and region.'],
    ]
    story.append(make_pdf_table(api_rows, [27*mm, 69*mm, 76*mm], styles))
    story += [Spacer(1, 7), P('<b>Validation / concurrency:</b> donation creation validates address, region, food type, quantity, units, expiry and coordinate ranges. NGO acceptance uses a conditional update to avoid accepting an already-claimed or expired request. Expired broadcasts are marked expired when request feeds are queried.', 'BodyX')]

    story += [PageBreak(), P('4. Database and Persistence', 'Section'), P('Prisma connects the application to SQLite through <b>DATABASE_URL</b>. The Prisma client singleton is in lib/prisma.ts; schema and migration SQL are under prisma/. The documented local database file is prisma/dev.db.', 'BodyX')]
    db_rows = [
        ['Model', 'Important fields / purpose'],
        ['Donation', 'Pickup address and optional coordinates; food type, quantity and unit; expiry; status; selected NGO; region; timestamps. Indexed by status and createdAt.'],
        ['NgoAccount', 'Name and unique normalized name key; scrypt password hash; region; active flag; availability; review status; creation date.'],
        ['NgoRequestDecline', 'Donation-to-NGO decline record; unique pair prevents duplicate declines. Relations cascade on deletion.'],
    ]
    story.append(make_pdf_table(db_rows, [35*mm, 137*mm], styles))
    story += [P('Data lifecycle', 'Subsection')]
    for text in ['Donation is created with status <b>broadcasting</b>.', 'An NGO can accept only if it is active/approved, available, within region (or Unassigned), and the donation is not expired or already accepted.', 'Accepted status transitions to <b>in_progress</b>, then <b>completed</b>. Overdue broadcasts become <b>expired</b>.', 'Declines are per NGO and do not hide the same request from every NGO.']: story.append(bullet(text))
    story += [P('Migration commands', 'Subsection'), P('<b>npm run db:migrate</b> applies Prisma development migrations. <b>npm run db:studio</b> opens Prisma Studio for local inspection. Migrations currently present are initial schema, NGO accounts, admin regions, availability/declines, and NGO review status.', 'BodyX')]

    story += [P('5. Chatbot and External Integrations', 'Section'), P('The floating assistant UI is components/help-chatbot.tsx. It posts the recent conversation to <b>POST /api/help</b>; the server uses <b>@google/genai</b> and GoogleGenAI to call the Gemini Developer API. The key is server-only and must not use a NEXT_PUBLIC_ prefix.', 'BodyX')]
    for text in ['Model selection uses GEMINI_MODEL when present; otherwise the route falls back to <b>gemini-3.8-flash</b> in the current source.', 'The route supplies a FoodRescue-specific system instruction, sends up to 12 messages, truncates message text, sets a maximum output token count, and caps requests at 15 per minute per forwarded client IP using in-memory state.', 'The assistant has no live web-search grounding in this configuration; answers may be stale. The client shows an error returned by the server when provider calls fail.', 'Reverse geocoding uses the public OpenStreetMap Nominatim endpoint with a User-Agent and an eight-second timeout.', 'Vercel Analytics is rendered only when NODE_ENV is production.']: story.append(bullet(text))

    story += [P('6. Authentication and Security', 'Section')]
    auth_rows = [
        ['Control', 'Implementation'],
        ['NGO passwords', 'Node crypto scrypt with per-password random salt; constant-time comparison.'],
        ['NGO session', 'HMAC-SHA256 signed payload, HTTP-only cookie, seven-day lifetime; active and approved account rechecked against database.'],
        ['Admin session', 'Environment-based username/password; HMAC-SHA256 signed cookie, HTTP-only, SameSite strict, eight-hour lifetime.'],
        ['Environment files', '.env and .env*.local are ignored by Git. Keep credentials there; share only placeholders.'],
        ['Production', 'Set a unique NGO_SESSION_SECRET of at least 32 random bytes and strong unique admin credentials. Use HTTPS and verify deployment secret management.'],
    ]
    story.append(make_pdf_table(auth_rows, [36*mm, 136*mm], styles))
    story += [Spacer(1, 7), P('<b>Secret incident note:</b> A real-looking Gemini key was found in .env.example during documentation generation and has been replaced with a placeholder. Revoke/rotate that key in Google AI Studio and review repository history if it was committed. This report does not print or embed the credential.', 'Callout')]

    story += [P('7. Environment and Local Setup', 'Section')]
    env_rows = [
        ['Variable', 'Purpose'],
        ['DATABASE_URL', 'Prisma connection string; sample SQLite value is file:./dev.db.'],
        ['NGO_SESSION_SECRET', 'Signs NGO and admin session payloads.'],
        ['ADMIN_USERNAME / ADMIN_PASSWORD', 'Credentials for the administrator portal.'],
        ['GEMINI_API_KEY', 'Private server-side Google Gemini API key.'],
        ['GEMINI_MODEL', 'Optional model override; current documented default is gemini-3.8-flash.'],
    ]
    story.append(make_pdf_table(env_rows, [49*mm, 123*mm], styles))
    story += [P('Typical local workflow', 'Subsection')]
    for text in ['Install dependencies using the package manager used by the project (npm install; package metadata also identifies pnpm).', 'Copy the safe template .env.example to .env and set local values. Never commit real credentials.', 'Run <b>npm run db:migrate</b>.', 'Start with <b>npm run dev</b>; open the local URL printed by Next.js.', 'Production: build with <b>npm run build</b> and start with <b>npm start</b>.']: story.append(bullet(text))

    story += [P('8. Known Limitations and Notes', 'Section')]
    for text in ['Prototype only: sample NGO entries/actions are not independently verified; no real pickup or food safety guarantee is provided.', 'SQLite and local dev database are configured; a production database/deployment topology is not specified in the reviewed project files.', 'The chatbot rate-limit map is process-local and resets on restart; distributed production deployments need shared rate limiting.', 'No live search for chatbot; validate time-sensitive answers independently.', 'next.config.mjs sets typescript.ignoreBuildErrors=true; remove or address this before relying on production build checks.', 'Review API access controls and data minimization before deployment, especially donation addresses and NGO personal data.']: story.append(bullet(text))
    story += [Spacer(1, 10), HRFlowable(width='100%', thickness=1, color=GOLD), Spacer(1, 6), P('Source of truth: project source, Prisma schema/migrations, package.json and README.md in this workspace. Generated 2026-10-09.', 'SmallX')]

    doc.build(story, onFirstPage=draw_pdf_footer, onLaterPages=draw_pdf_footer)


def make_pdf_table(rows, widths, styles):
    converted = []
    for row_index, row in enumerate(rows):
        style_name = 'CellHead' if row_index == 0 else 'CellX'
        converted.append([Paragraph(str(value), styles[style_name]) for value in row])
    table = Table(converted, colWidths=widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),NAVY),
        ('BACKGROUND',(0,1),(-1,-1),colors.white),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white, PALE]),
        ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#CCD7DE')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),
        ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),
    ]))
    return table


def draw_pdf_footer(canvas, doc):
    canvas.saveState()
    w, _ = A4
    canvas.setStrokeColor(colors.HexColor('#DCE5E8'))
    canvas.line(18*mm, 13*mm, w-18*mm, 13*mm)
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(18*mm, 8*mm, 'FoodRescue · Project Overview · No secret values included')
    canvas.drawRightString(w-18*mm, 8*mm, f'{doc.page}')
    canvas.restoreState()


# ----- PowerPoint generation -----
SLIDE_BG = RGBColor(16, 36, 58)
SLIDE_PANEL = RGBColor(27, 51, 73)
SLIDE_GREEN = RGBColor(63, 190, 137)
SLIDE_GOLD = RGBColor(243, 182, 31)
SLIDE_WHITE = RGBColor(247, 250, 252)
SLIDE_MUTED = RGBColor(184, 199, 211)


def add_text(slide, x, y, w, h, text, size=18, color=SLIDE_WHITE, bold=False, font='Aptos', align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_base(prs, title, section, slide_no):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = SLIDE_BG
    add_text(slide, 0.55, 0.28, 4.2, 0.3, 'FOODRESCUE  /  PROJECT BRIEF', 9, SLIDE_GREEN, True)
    add_text(slide, 0.55, 0.70, 12.1, 0.55, title, 26, SLIDE_WHITE, True)
    add_text(slide, 0.58, 1.27, 11.8, 0.30, section, 10, SLIDE_MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.68), Inches(12.2), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = SLIDE_GOLD; line.line.fill.background()
    add_text(slide, 12.0, 7.13, 0.7, 0.22, f'{slide_no:02}', 9, SLIDE_MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_card(slide, x, y, w, h, title, body, accent=SLIDE_GREEN, body_size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = SLIDE_PANEL
    shape.line.color.rgb = RGBColor(54, 82, 103)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, x+0.20, y+0.16, w-0.38, 0.35, title, 15, accent, True)
    add_text(slide, x+0.20, y+0.60, w-0.38, h-0.72, body, body_size, SLIDE_WHITE)


def add_bullets(slide, x, y, w, items, size=16, gap=0.70):
    for i, item in enumerate(items):
        yy = y + i*gap
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy+0.10), Inches(0.11), Inches(0.11))
        dot.fill.solid(); dot.fill.fore_color.rgb = SLIDE_GREEN; dot.line.fill.background()
        add_text(slide, x+0.25, yy, w-0.25, gap-0.06, item, size, SLIDE_WHITE)


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = 'FoodRescue Project Overview'
    prs.core_properties.subject = 'Application architecture, frontend, backend, database and Gemini chatbot'
    prs.core_properties.author = 'FoodRescue Project'

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = SLIDE_BG
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = SLIDE_GOLD; stripe.line.fill.background()
    add_text(slide, 0.8, 1.0, 5.2, 0.35, 'FOODRESCUE  /  PROJECT OVERVIEW', 11, SLIDE_GREEN, True)
    add_text(slide, 0.8, 1.75, 9.5, 1.0, 'Surplus food,\nconnected to community.', 34, SLIDE_WHITE, True)
    add_text(slide, 0.82, 4.0, 8.3, 0.7, 'Technical guide · frontend · backend · database · AI assistant', 18, SLIDE_MUTED)
    add_text(slide, 0.82, 6.58, 7.5, 0.3, 'Project prototype · Generated 2026-10-09', 10, SLIDE_MUTED)
    add_card(slide, 9.1, 1.75, 3.2, 1.25, 'FRONTEND', 'Next.js · React · TypeScript', SLIDE_GREEN, 13)
    add_card(slide, 9.1, 3.25, 3.2, 1.25, 'DATA', 'Prisma · SQLite', SLIDE_GOLD, 13)
    add_card(slide, 9.1, 4.75, 3.2, 1.25, 'CHAT', 'Google Gemini Developer API', SLIDE_GREEN, 13)

    slide = add_base(prs, 'What FoodRescue does', 'A prototype workflow for donors, NGO operators and administrators.', 2)
    add_card(slide, 0.7, 2.05, 3.8, 3.5, '01  DONATE', 'Provide pickup location, food type, quantity and expiry. Review NGO options and broadcast a request.', SLIDE_GREEN, 16)
    add_card(slide, 4.75, 2.05, 3.8, 3.5, '02  RESPOND', 'Approved NGOs see eligible requests in their region, accept or decline, and update pickup progress.', SLIDE_GOLD, 16)
    add_card(slide, 8.8, 2.05, 3.8, 3.5, '03  OPERATE', 'Administrators review NGO accounts, assign regions, manage access and view operational summaries.', SLIDE_GREEN, 16)
    add_text(slide, 0.75, 6.0, 11.7, 0.45, 'Prototype note: sample NGO listings/actions are not independently verified.', 13, SLIDE_MUTED)

    slide = add_base(prs, 'Frontend: pages and experience', 'Next.js App Router · React 19 · TypeScript · Tailwind CSS 4', 3)
    add_bullets(slide, 0.9, 2.0, 7.2, [
        'Home page: donor entry, hero and impact sections.',
        'NGO selection: donation details and available organizations.',
        'NGO portal: registration/login, request feed and status updates.',
        'Admin portal: protected sign-in and operations dashboard.',
        'Global floating help chatbot and shared site navigation.',
    ], 16, 0.78)
    add_card(slide, 8.7, 2.1, 3.6, 3.6, 'KEY UI MODULES', 'hero-section\nngo-selection\nngo-dashboard\nadmin-dashboard\nhelp-chatbot\nsite-header', SLIDE_GOLD, 15)
    add_text(slide, 0.9, 6.45, 11.6, 0.3, 'Key source areas: app/page.tsx · app/select-ngo · app/ngo · app/admin · components/', 11, SLIDE_MUTED)

    slide = add_base(prs, 'Backend: API and request lifecycle', 'Route handlers validate input and use Prisma for persistence.', 4)
    add_card(slide, 0.7, 2.0, 3.8, 3.7, 'DONATION APIs', 'POST /api/donations\nGET /api/donations\nDonation item/status routes\n\nValidates location, quantity, unit, expiry and coordinates.', SLIDE_GREEN, 14)
    add_card(slide, 4.75, 2.0, 3.8, 3.7, 'NGO APIs', 'NGO auth and availability\nRequest feed\nAccept / decline\nPickup status progression\n\nRegion, availability and expiry eligibility checks.', SLIDE_GOLD, 14)
    add_card(slide, 8.8, 2.0, 3.8, 3.7, 'ADMIN + SERVICES', 'Admin auth/dashboard/NGO management\nPOST /api/help → Gemini\nGET /api/location/reverse → Nominatim', SLIDE_GREEN, 14)
    add_text(slide, 0.8, 6.1, 11.7, 0.4, 'Acceptance uses conditional database updates to reduce double-claim races.', 13, SLIDE_MUTED)

    slide = add_base(prs, 'Database and persistence', 'Prisma ORM · SQLite local database · versioned migrations', 5)
    add_card(slide, 0.75, 2.05, 3.8, 2.3, 'Donation', 'Pickup + optional coordinates\nFood + quantity + expiry\nRegion + status + selected NGO', SLIDE_GREEN, 14)
    add_card(slide, 4.75, 2.05, 3.8, 2.3, 'NgoAccount', 'Name + unique key\nScrypt password hash\nRegion + active/review/availability', SLIDE_GOLD, 14)
    add_card(slide, 8.75, 2.05, 3.8, 2.3, 'NgoRequestDecline', 'Donation ↔ NGO decline record\nUnique pair prevents duplicate declines\nCascading relations', SLIDE_GREEN, 14)
    add_text(slide, 0.9, 4.9, 11.5, 0.8, 'Connection: DATABASE_URL → Prisma datasource (SQLite). Local database documented at prisma/dev.db. Schema: prisma/schema.prisma; migrations: prisma/migrations/.', 17, SLIDE_WHITE)
    add_text(slide, 0.9, 6.05, 11.5, 0.45, 'Commands: npm run db:migrate · npm run db:studio', 13, SLIDE_MUTED)

    slide = add_base(prs, 'Chatbot: Gemini integration', 'Browser chat UI → Next.js API route → Google Gemini Developer API', 6)
    add_card(slide, 0.75, 2.0, 3.7, 3.8, 'CLIENT', 'components/help-chatbot.tsx\n\nSends recent conversation to POST /api/help. Key is never sent from the browser.', SLIDE_GREEN, 15)
    add_card(slide, 4.8, 2.0, 3.7, 3.8, 'SERVER', 'app/api/help/route.ts\n\nUses @google/genai. Reads GEMINI_API_KEY. Model override via GEMINI_MODEL; default in source: gemini-3.8-flash.', SLIDE_GOLD, 14)
    add_card(slide, 8.85, 2.0, 3.7, 3.8, 'GUARDRAILS', 'FoodRescue system instruction\nMessage and token limits\nIn-memory rate limit\nNo live web search grounding', SLIDE_GREEN, 14)
    add_text(slide, 0.9, 6.2, 11.4, 0.45, 'Keep the key in server-side .env / deployment secrets; never use NEXT_PUBLIC_GEMINI_API_KEY.', 14, SLIDE_MUTED)

    slide = add_base(prs, 'Authentication and security', 'Signed HTTP-only sessions; NGO credentials verified against database.', 7)
    add_bullets(slide, 0.9, 2.0, 7.2, [
        'NGO passwords: Node.js scrypt with random salts.',
        'NGO sessions: HMAC-SHA256 signed; seven-day expiry.',
        'Admin credentials: environment variables; signed session, eight-hour expiry.',
        'Cookies are HTTP-only; secure flag enabled in production.',
        'Production needs a unique 32+ byte session secret and strong admin password.',
    ], 15, 0.76)
    add_card(slide, 8.7, 2.15, 3.7, 3.5, 'CREDENTIAL HYGIENE', 'Use ignored .env or .env.local for real keys.\n\nShare .env.example with placeholders only.\n\nRotate the Gemini key previously found in the sample file.', SLIDE_GOLD, 14)

    slide = add_base(prs, 'External services and data flow', 'Supporting services used by the current source.', 8)
    add_card(slide, 0.8, 2.0, 3.7, 3.3, 'OPENSTREETMAP NOMINATIM', 'Reverse geocoding from latitude/longitude to address and region. Server-side request includes User-Agent and timeout.', SLIDE_GREEN, 15)
    add_card(slide, 4.8, 2.0, 3.7, 3.3, 'GEMINI API', 'Generates conversational assistant responses. No live search is configured; users should verify time-sensitive claims.', SLIDE_GOLD, 15)
    add_card(slide, 8.8, 2.0, 3.7, 3.3, 'VERCEL ANALYTICS', 'Analytics component is included only when NODE_ENV is production.', SLIDE_GREEN, 15)
    add_text(slide, 1.0, 5.85, 11.2, 0.55, 'Donor browser → Next.js UI → API route → Prisma/SQLite or external provider → JSON response → UI', 17, SLIDE_WHITE, True, align=PP_ALIGN.CENTER)

    slide = add_base(prs, 'Run and configure locally', 'Environment configuration and common npm scripts.', 9)
    add_card(slide, 0.75, 2.0, 5.5, 3.9, 'ENVIRONMENT VARIABLES', 'DATABASE_URL — SQLite connection\nNGO_SESSION_SECRET — signs sessions\nADMIN_USERNAME / ADMIN_PASSWORD\nGEMINI_API_KEY — private server key\nGEMINI_MODEL — optional model override', SLIDE_GREEN, 15)
    add_card(slide, 6.55, 2.0, 5.6, 3.9, 'DEVELOPER WORKFLOW', '1  Install dependencies\n2  Copy .env.example to .env; set local values\n3  npm run db:migrate\n4  npm run dev\n5  npm run build / npm start', SLIDE_GOLD, 15)
    add_text(slide, 0.85, 6.35, 11.7, 0.34, 'Package metadata specifies pnpm; README documents npm commands. Use the package manager/lockfile chosen for this checkout consistently.', 11, SLIDE_MUTED)

    slide = add_base(prs, 'Current limitations and next steps', 'Important before treating the prototype as a production service.', 10)
    add_bullets(slide, 0.85, 1.98, 11.8, [
        'Sample NGO data and actions are not independently verified; no real pickup or food-safety guarantee.',
        'SQLite is configured locally; production database and deployment topology need explicit planning.',
        'Chat rate limits are in process memory; production scale needs a shared limiter and monitoring.',
        'No live web search for AI answers; evaluate data privacy, provider terms, costs and user disclosures.',
        'next.config.mjs ignores TypeScript build errors; re-enable checks and resolve diagnostics before production.',
        'Rotate exposed Gemini key; audit Git history and keep all secrets out of generated/shareable assets.',
    ], 15, 0.71)
    add_text(slide, 0.9, 6.45, 11.4, 0.35, 'Source of truth: workspace source, schema/migrations, package.json and README.md · 2026-10-09', 10, SLIDE_MUTED)

    prs.save(PPT_PATH)


if __name__ == '__main__':
    build_pdf()
    build_ppt()
    print(f'Created {PDF_PATH.name} ({PDF_PATH.stat().st_size:,} bytes)')
    print(f'Created {PPT_PATH.name} ({PPT_PATH.stat().st_size:,} bytes)')
